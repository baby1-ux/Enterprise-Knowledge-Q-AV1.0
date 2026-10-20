"""DOCX 解析：提取段落与表格文本。"""

from __future__ import annotations

from pathlib import Path

from ...core.logging import get_logger
from .base import BaseParser, ParseError, ParsedBlock

logger = get_logger("parser.docx")


class DocxParser(BaseParser):
    """基于 python-docx 的 Word 解析器。"""

    extensions = (".docx",)

    def parse(self, path: Path) -> list[ParsedBlock]:
        self._ensure_exists(path)
        try:
            from docx import Document as DocxDocument
        except ImportError as exc:  # pragma: no cover
            raise ParseError("缺少 python-docx 依赖，无法解析 Word 文件") from exc

        try:
            document = DocxDocument(str(path))
        except Exception as exc:
            raise ParseError(f"Word 文件无法打开：{exc}") from exc

        parts: list[str] = []
        for paragraph in document.paragraphs:
            text = paragraph.text.strip()
            if text:
                parts.append(text)

        for table in document.tables:
            for row in table.rows:
                cells = [cell.text.strip().replace("\n", " ") for cell in row.cells]
                line = " | ".join(cell for cell in cells if cell)
                if line:
                    parts.append(line)

        text = self._clean_text("\n".join(parts))
        if not text:
            raise ParseError("未从 Word 文档中提取到文本内容")
        logger.info("Word 解析完成：%s", path.name)
        return [ParsedBlock(text=text, page_no=None, meta={"source": path.name})]
