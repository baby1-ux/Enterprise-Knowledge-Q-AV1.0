"""PPT / PPTX 解析：按幻灯片提取文本框与表格文本。"""

from __future__ import annotations

from pathlib import Path

from ...core.logging import get_logger
from .base import BaseParser, ParseError, ParsedBlock

logger = get_logger("parser.pptx")


class PptxParser(BaseParser):
    """基于 python-pptx 的演示文稿解析器。"""

    extensions = (".pptx",)

    def parse(self, path: Path) -> list[ParsedBlock]:
        self._ensure_exists(path)
        try:
            from pptx import Presentation
        except ImportError as exc:  # pragma: no cover
            raise ParseError("缺少 python-pptx 依赖，无法解析 PPT 文件") from exc

        try:
            presentation = Presentation(str(path))
        except Exception as exc:
            raise ParseError(f"PPT 文件无法打开：{exc}") from exc

        blocks: list[ParsedBlock] = []
        for index, slide in enumerate(presentation.slides):
            parts: list[str] = []
            for shape in slide.shapes:
                parts.extend(self._extract_shape_text(shape))
            text = self._clean_text("\n".join(part for part in parts if part))
            if text:
                blocks.append(ParsedBlock(text=text, page_no=index + 1, meta={"source": path.name}))

        if not blocks:
            raise ParseError("未从 PPT 中提取到文本内容")
        logger.info("PPT 解析完成：%s，共 %s 页有文本", path.name, len(blocks))
        return blocks

    def _extract_shape_text(self, shape) -> list[str]:
        """递归提取形状中的文字（含组合形状与表格）。"""
        texts: list[str] = []

        if getattr(shape, "has_text_frame", False):
            for paragraph in shape.text_frame.paragraphs:
                line = "".join(run.text for run in paragraph.runs).strip()
                if line:
                    texts.append(line)

        if getattr(shape, "has_table", False):
            for row in shape.table.rows:
                cells = [cell.text.strip().replace("\n", " ") for cell in row.cells]
                line = " | ".join(cell for cell in cells if cell)
                if line:
                    texts.append(line)

        if getattr(shape, "shape_type", None) == 6 and hasattr(shape, "shapes"):  # 组合形状
            for child in shape.shapes:
                texts.extend(self._extract_shape_text(child))

        return texts


class PptParser(BaseParser):
    """旧版 .ppt 文件的降级处理：提示用户另存为 pptx。"""

    extensions = (".ppt",)

    def parse(self, path: Path) -> list[ParsedBlock]:
        raise ParseError("暂不支持旧版 .ppt 格式，请在 Office 中另存为 .pptx 后再上传")
