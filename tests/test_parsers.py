"""文档解析单元测试：Markdown / TXT / 不支持类型，以及可选的 PPT、PDF。"""

from __future__ import annotations

import importlib.util
from pathlib import Path

import pytest

from app.services.parsing.base import ParsedBlock, ParseError
from app.services.parsing.factory import get_parser, parse_document, supported_extensions

HAS_PPTX = importlib.util.find_spec("pptx") is not None
HAS_PYPDF = importlib.util.find_spec("pypdf") is not None


def test_supported_extensions_cover_required_types() -> None:
    """任务要求的文件类型都要被支持。"""
    extensions = set(supported_extensions())
    for required in (".pdf", ".pptx", ".md", ".png", ".jpg"):
        assert required in extensions


def test_markdown_parser_splits_by_heading(tmp_path: Path) -> None:
    """Markdown 按标题切分为多个逻辑块。"""
    file_path = tmp_path / "手册.md"
    file_path.write_text(
        "# 产品保修政策\n\n## 保修期限\n\n公司提供一年免费保修服务。\n\n## 保修流程\n\n在官网提交申请。\n",
        encoding="utf-8",
    )
    blocks = parse_document(file_path)
    assert len(blocks) >= 2
    assert any("一年免费保修" in block.text for block in blocks)
    assert any(block.meta.get("title") for block in blocks)


def test_markdown_parser_rejects_empty_file(tmp_path: Path) -> None:
    """空 Markdown 文件报解析错误。"""
    file_path = tmp_path / "空文件.md"
    file_path.write_text("   \n\n", encoding="utf-8")
    with pytest.raises(ParseError):
        parse_document(file_path)


def test_text_parser_splits_paragraphs(tmp_path: Path) -> None:
    """纯文本按空行切分为段落块。"""
    file_path = tmp_path / "说明.txt"
    file_path.write_text("第一段内容。\n\n第二段内容。\n", encoding="utf-8")
    blocks = parse_document(file_path)
    assert len(blocks) == 2
    assert blocks[0].text == "第一段内容。"


def test_text_parser_gbk_encoding(tmp_path: Path) -> None:
    """非 UTF-8 编码的文本文件也能正常读取。"""
    file_path = tmp_path / "gbk.txt"
    file_path.write_bytes("中文内容测试。".encode("gbk"))
    blocks = parse_document(file_path)
    assert "中文内容测试" in blocks[0].text


def test_unsupported_extension_raises(tmp_path: Path) -> None:
    """不支持的文件类型给出中文提示。"""
    file_path = tmp_path / "数据.xyz"
    file_path.write_text("内容", encoding="utf-8")
    with pytest.raises(ParseError) as excinfo:
        get_parser(file_path.name)
    assert "暂不支持" in str(excinfo.value)


def test_missing_file_raises(tmp_path: Path) -> None:
    """文件不存在时报错。"""
    with pytest.raises(ParseError):
        parse_document(tmp_path / "不存在.md")


def test_markdown_blocks_keep_source_meta(tmp_path: Path) -> None:
    """解析块携带来源文件名。"""
    file_path = tmp_path / "制度.md"
    file_path.write_text("# 考勤制度\n\n每天 9 点上班。\n", encoding="utf-8")
    blocks = parse_document(file_path)
    assert all(block.meta.get("source") == "制度.md" for block in blocks)


@pytest.mark.skipif(not HAS_PPTX, reason="未安装 python-pptx")
def test_pptx_parser_extracts_slide_text(tmp_path: Path) -> None:
    """PPT 解析提取幻灯片文本与页码。"""
    from pptx import Presentation
    from pptx.util import Inches

    presentation = Presentation()
    slide = presentation.slides.add_slide(presentation.slide_layouts[6])
    textbox = slide.shapes.add_textbox(Inches(1), Inches(1), Inches(4), Inches(1))
    textbox.text_frame.text = "第二季度销售目标为增长百分之二十"
    second = presentation.slides.add_slide(presentation.slide_layouts[6])
    textbox2 = second.shapes.add_textbox(Inches(1), Inches(1), Inches(4), Inches(1))
    textbox2.text_frame.text = "重点客户包括华东区代理商"

    file_path = tmp_path / "汇报.pptx"
    presentation.save(str(file_path))

    blocks = parse_document(file_path)
    assert len(blocks) == 2
    assert "销售目标" in blocks[0].text
    assert blocks[0].page_no == 1
    assert blocks[1].page_no == 2


@pytest.mark.skipif(not HAS_PYPDF, reason="未安装 pypdf")
def test_pdf_parser_reports_broken_file(tmp_path: Path) -> None:
    """损坏的 PDF 文件给出中文错误提示。"""
    file_path = tmp_path / "坏文件.pdf"
    file_path.write_bytes(b"%PDF-1.4 this is not a real pdf")
    with pytest.raises(ParseError):
        parse_document(file_path)


def test_old_ppt_format_is_rejected(tmp_path: Path) -> None:
    """旧版 .ppt 给出明确提示。"""
    file_path = tmp_path / "旧格式.ppt"
    file_path.write_bytes(b"not a real ppt")
    with pytest.raises(ParseError) as excinfo:
        parse_document(file_path)
    assert "pptx" in str(excinfo.value)


def test_parsed_block_is_empty() -> None:
    """空块判断。"""
    assert ParsedBlock(text="   ").is_empty() is True
    assert ParsedBlock(text="内容").is_empty() is False
